# -*- coding: utf-8 -*-
"""输出完成表：复制样表 -> 从指定行开始写入转换结果。"""
from __future__ import annotations

import os
import shutil
from copy import copy
from dataclasses import dataclass

import openpyxl
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from .excel_reader import cell_text


@dataclass
class CellFormat:
    """输出到 Excel 的字体/颜色设置。默认不启用，完全跟随样表原格式。"""

    enabled: bool = False
    family: str = "微软雅黑"
    size: float = 11
    color: str = "#000000"
    bold: bool = False
    align: str = ""              # "" | left | center | right
    apply_header: bool = False

    @staticmethod
    def from_settings(s) -> "CellFormat":
        return CellFormat(
            enabled=s.get_bool("fmt.enabled"),
            family=s.get_str("fmt.family") or "微软雅黑",
            size=float(s.get("fmt.size") or 11),
            color=s.get_str("fmt.color") or "#000000",
            bold=s.get_bool("fmt.bold"),
            align=s.get_str("fmt.align"),
            apply_header=s.get_bool("fmt.apply_header"),
        )


def argb(color: str) -> str:
    """'#RRGGBB' / 'RRGGBB' / 颜色名 -> openpyxl 认的颜色串。"""
    c = (color or "").strip()
    if not c:
        return "FF000000"
    if c.startswith("#"):
        c = c[1:]
    if len(c) == 6:
        return "FF" + c.upper()
    if len(c) == 8:
        return c.upper()
    return c


def apply_format(cell, fmt: CellFormat):
    """在已有（从样表复制来的）样式基础上，覆盖字体与对齐。边框、数字格式不动。"""
    if not fmt.enabled:
        return
    try:
        cell.font = Font(
            name=fmt.family,
            size=fmt.size,
            bold=fmt.bold,
            color=argb(fmt.color),
        )
    except Exception:
        pass
    if fmt.align in ("left", "center", "right"):
        try:
            cell.alignment = Alignment(horizontal=fmt.align, vertical="center")
        except Exception:
            pass


def export(
    sample_path: str,
    output_path: str,
    sheet: str,
    spec,
    rows: list,
    clear_example_rows: bool = False,
    fmt: CellFormat | None = None,
) -> dict:
    """把结果写进样表的副本。

    - 复制一份样表到 output_path（保留原样表的表头、格式、说明文字）
    - 从 spec.data_start_row 行开始，按样表列顺序写入每一行数据
    - clear_example_rows=True 时清空第 (header_row+1) 行的示例数据
    - fmt 启用时，把数据行（可选含表头行）的字体/字号/颜色/对齐改成指定的
    """
    fmt = fmt or CellFormat()
    out_dir = os.path.dirname(os.path.abspath(output_path))
    if out_dir:
        os.makedirs(out_dir, exist_ok=True)
    if os.path.abspath(sample_path) == os.path.abspath(output_path):
        raise ValueError("输出文件不能和样表是同一个文件，请换个名字")

    shutil.copy2(sample_path, output_path)

    wb = openpyxl.load_workbook(output_path)
    if sheet not in wb.sheetnames:
        sheet = wb.sheetnames[0]
    ws = wb[sheet]

    header_row = spec.header_row
    start_row = spec.data_start_row
    columns = [c for c in spec.columns if c in spec.col_index]

    # 1) 需要时清空示例数据行
    if clear_example_rows:
        ex = header_row + 1
        if ex < start_row:
            for c in range(1, ws.max_column + 1):
                ws.cell(ex, c).value = None

    # 2) 清空起始行以下的旧内容
    for r in range(start_row, ws.max_row + 1):
        for c in range(1, ws.max_column + 1):
            ws.cell(r, c).value = None

    # 3) 用示例行的样式作为模板
    template_styles = {}
    ex_row = header_row + 1
    if ex_row < start_row:
        for c in range(1, ws.max_column + 1):
            template_styles[c] = copy(ws.cell(ex_row, c)._style)

    # 4) 写数据
    written = 0
    for i, row in enumerate(rows):
        r = start_row + i
        empty = True
        for name in columns:
            col = spec.col_index[name]
            v = row.get(name)
            if isinstance(v, str):
                v = v.strip()
            cell = ws.cell(r, col)
            if v is None or v == "":
                cell.value = None
            else:
                cell.value = v
                empty = False
        if not empty:
            written += 1
        if template_styles:
            for c, style in template_styles.items():
                try:
                    ws.cell(r, c)._style = copy(style)
                except Exception:
                    pass
        if fmt.enabled:
            for c in range(1, ws.max_column + 1):
                apply_format(ws.cell(r, c), fmt)

    # 5) 表头行也改格式（可选）
    if fmt.enabled and fmt.apply_header:
        for c in range(1, ws.max_column + 1):
            apply_format(ws.cell(header_row, c), fmt)

    wb.save(output_path)
    return {
        "输出文件": output_path,
        "工作表": sheet,
        "起始行": start_row,
        "写入行数": written,
        "列数": len(columns),
        "自定义格式": "已应用" if fmt.enabled else "跟随样表",
    }


def column_letters(spec, names) -> list:
    return [get_column_letter(spec.col_index[n]) for n in names if n in spec.col_index]
