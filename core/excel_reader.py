# -*- coding: utf-8 -*-
"""读取初始表（申报汇总表）与样表，自动识别表头结构。"""
from __future__ import annotations

import re
from dataclasses import dataclass, field

import openpyxl

NAME_MARKER = "姓名"
INCOME_FIELD = "所得项目"
PERIOD_RE = re.compile(r"^\s*(\d{4})\s*年\s*(\d{1,2})\s*月\s*$")


def cell_text(v) -> str:
    """把单元格值规范成字符串（去掉首尾空白，整数不显示 .0）。"""
    if v is None:
        return ""
    if isinstance(v, bool):
        return str(v)
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def read_workbook_rows(path: str, sheet: str | None = None):
    """读取整张表为二维列表，返回 (工作表名列表, 实际工作表名, 行数据)。"""
    wb = openpyxl.load_workbook(path, data_only=True)
    names = list(wb.sheetnames)
    if not names:
        raise ValueError("工作簿里没有任何工作表")
    if sheet is None or sheet not in names:
        sheet = names[0]
    ws = wb[sheet]
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    max_cols = max((len(r) for r in rows), default=0)
    for r in rows:
        if len(r) < max_cols:
            r.extend([None] * (max_cols - len(r)))
    wb.close()
    return names, sheet, rows


@dataclass
class SourceTable:
    """初始表解析结果。"""

    path: str = ""
    sheet: str = ""
    mode: str = "flat"                 # block=按人分块；flat=一行一条流水表
    fields: list = field(default_factory=list)     # 可参与规则的字段名
    records: list = field(default_factory=list)    # 每条记录 dict{字段名: 值}
    info: dict = field(default_factory=dict)       # 附加信息（人数、字段数等）
    preview: list = field(default_factory=list)    # 前若干行原始数据，用于展示

    @property
    def field_values(self) -> dict:
        """统计每个字段出现过的不同取值（用于筛选值下拉）。"""
        out: dict[str, list] = {}
        for f in self.fields:
            seen, vals = set(), []
            for rec in self.records:
                t = cell_text(rec.get(f))
                if t and t not in seen:
                    seen.add(t)
                    vals.append(t)
                if len(vals) >= 200:
                    break
            out[f] = vals
        return out


def parse_source(path: str, sheet: str | None = None) -> SourceTable:
    """解析初始表。自动识别「按人分块」结构或普通流水表结构。"""
    _, sheet, rows = read_workbook_rows(path, sheet)
    st = SourceTable(path=path, sheet=sheet)
    st.preview = rows[:40]

    markers = [i for i, r in enumerate(rows) if cell_text(r[0]) == NAME_MARKER]
    if markers:
        _parse_block(rows, markers, st)
    else:
        _parse_flat(rows, st)

    st.info["记录数"] = len(st.records)
    st.info["字段数"] = len(st.fields)
    return st


def _looks_like_header_row(row) -> bool:
    return any(cell_text(v) == INCOME_FIELD for v in row)


def _parse_block(rows, markers, st: SourceTable) -> None:
    """按人分块：姓名行 -> 人员信息行 -> 字段表头行 -> 若干数据行。"""
    st.mode = "block"

    first = markers[0]
    # 人员信息字段（姓名/国籍/证件类型/证件号码…）来自姓名行的各列标题
    info_cols = {c: cell_text(v) for c, v in enumerate(rows[first]) if cell_text(v)}

    # 字段表头行：姓名行之后、包含「所得项目」的那一行
    def find_header(m):
        for i in range(m + 1, min(m + 4, len(rows))):
            if _looks_like_header_row(rows[i]):
                return i
        cand = range(m + 1, min(m + 3, len(rows)))
        return max(cand, key=lambda i: sum(1 for v in rows[i] if cell_text(v)))

    hdr = find_header(first)
    data_cols = {c: cell_text(v) for c, v in enumerate(rows[hdr]) if cell_text(v)}

    fields: list[str] = []
    for c in sorted(info_cols):
        if info_cols[c] not in fields:
            fields.append(info_cols[c])
    for c in sorted(data_cols):
        if data_cols[c] not in fields:
            fields.append(data_cols[c])
    st.fields = fields

    for k, m in enumerate(markers):
        block_end = markers[k + 1] if k + 1 < len(markers) else len(rows)
        h = find_header(m)
        info_row = rows[m + 1] if m + 1 < len(rows) else []
        base = {nm: (info_row[c] if c < len(info_row) else None) for c, nm in info_cols.items()}
        for r in range(h + 1, block_end):
            row = rows[r]
            if not any(cell_text(v) for v in row):
                continue
            if cell_text(row[0]) == NAME_MARKER:
                continue
            rec = dict(base)
            for c, nm in data_cols.items():
                rec[nm] = row[c] if c < len(row) else None
            st.records.append(rec)


def _parse_flat(rows, st: SourceTable) -> None:
    """普通流水表：第一行有效表头 + 后续数据行。"""
    st.mode = "flat"
    if not rows:
        return
    limit = min(10, len(rows))
    hdr = max(range(limit), key=lambda i: sum(1 for v in rows[i] if cell_text(v)))
    cols = {c: cell_text(v) for c, v in enumerate(rows[hdr]) if cell_text(v)}
    st.fields = [cols[c] for c in sorted(cols)]
    for r in range(hdr + 1, len(rows)):
        row = rows[r]
        if not any(cell_text(v) for v in row):
            continue
        st.records.append({nm: (row[c] if c < len(row) else None) for c, nm in cols.items()})


@dataclass
class SampleSpec:
    """样表解析结果。"""

    path: str = ""
    sheet: str = ""
    header_row: int = 1
    data_start_row: int = 4
    columns: list = field(default_factory=list)      # 样表列名（按列顺序）
    col_index: dict = field(default_factory=dict)    # 列名 -> 1 基列号
    rename_rule_text: str = ""
    sheets: list = field(default_factory=list)


def parse_sample(
    path: str,
    sheet: str | None = None,
    header_row: int = 1,
    data_start_row: int = 4,
) -> SampleSpec:
    """解析样表：读取表头列、数据起始行，并提取样表中的重名规则说明。"""
    names, sheet, rows = read_workbook_rows(path, sheet)
    spec = SampleSpec(
        path=path,
        sheet=sheet,
        header_row=header_row,
        data_start_row=data_start_row,
        sheets=names,
    )
    hdr = rows[header_row - 1] if header_row - 1 < len(rows) else []
    for c, v in enumerate(hdr):
        nm = cell_text(v)
        if nm and nm not in spec.col_index:
            spec.columns.append(nm)
            spec.col_index[nm] = c + 1

    for row in rows:
        hit = ""
        for v in row:
            t = cell_text(v)
            if "重名" in t and len(t) > len(hit):
                hit = t
        if hit:
            spec.rename_rule_text = hit
            break
    return spec


def sample_preview_rows(path: str, sheet: str, header_row: int, data_start_row: int):
    """返回样表原始行（供界面预览）。"""
    _, _, rows = read_workbook_rows(path, sheet)
    return rows, header_row, data_start_row
